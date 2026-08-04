import io

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from app.extraction.normalize import Deck, DeckExtractionError, Slide


def extract_pptx(file_bytes: bytes, filename: str) -> Deck:
    try:
        prs = Presentation(io.BytesIO(file_bytes))
    except PackageNotFoundError as e:
        raise DeckExtractionError(
            f"Could not open '{filename}' as a PPTX. The file may be corrupted "
            f"or not a valid PowerPoint file."
        ) from e
    except Exception as e:
        raise DeckExtractionError(
            f"Unexpected error reading '{filename}' as a PPTX file."
        ) from e

    slide_list = list(prs.slides)
    if not slide_list:
        raise DeckExtractionError(f"'{filename}' has no slides.")

    slides: list[Slide] = []
    for i, slide in enumerate(slide_list):
        text_blocks: list[str] = []
        for shape in slide.shapes:
            text_blocks.extend(_extract_shape_text(shape))

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text

        slides.append(Slide(index=i + 1, text_blocks=text_blocks, notes=notes))

    deck = Deck(filename=filename, source_format="pptx", slides=slides)

    if deck.non_empty_slide_count == 0:
        raise DeckExtractionError(
            f"'{filename}' contains no extractable text on any slide. "
            f"If this deck is mostly images/screenshots, text extraction "
            f"won't pick up its content - OCR support is not available yet."
        )

    return deck


def _extract_shape_text(shape) -> list[str]:
    blocks: list[str] = []

    if shape.has_text_frame and shape.text_frame.text.strip():
        blocks.append(shape.text_frame.text)

    if shape.has_table:
        for row in shape.table.rows:
            row_text = " | ".join(cell.text for cell in row.cells if cell.text.strip())
            if row_text.strip():
                blocks.append(row_text)

    if shape.has_chart:
        chart = shape.chart
        title = ""
        try:
            if chart.has_title:
                title = chart.chart_title.text_frame.text
        except Exception:
            title = ""
        if title.strip():
            blocks.append(f"[Chart: {title.strip()}]")
        else:
            blocks.append("[Chart present, no title/caption extracted]")

    if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
        blocks.append("[Image present, no caption extracted]")

    if getattr(shape, "shapes", None) is not None:
        # Grouped shape - recurse into children
        for child in shape.shapes:
            blocks.extend(_extract_shape_text(child))

    return blocks
