import io

import fitz
import pytest
from pptx import Presentation

from app.extraction.normalize import DeckExtractionError
from app.extraction.pdf import extract_pdf
from app.extraction.pptx import extract_pptx


def make_pdf_bytes(pages_text: list[str]) -> bytes:
    doc = fitz.open()
    for text in pages_text:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pptx_bytes(slide_texts: list[str]) -> bytes:
    prs = Presentation()
    layout = prs.slide_layouts[1]
    for text in slide_texts:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_extract_pdf_happy_path():
    pdf_bytes = make_pdf_bytes(
        [
            "Problem: startups struggle to know if they're investment ready.",
            "Solution: an AI-powered readiness score.",
        ]
    )
    deck = extract_pdf(pdf_bytes, "deck.pdf")

    assert deck.source_format == "pdf"
    assert deck.slide_count == 2
    assert "investment ready" in deck.slides[0].text
    assert "AI-powered readiness score" in deck.slides[1].text
    assert deck.non_empty_slide_count == 2


def test_extract_pdf_rejects_corrupted_file():
    with pytest.raises(DeckExtractionError):
        extract_pdf(b"not a real pdf", "broken.pdf")


def test_extract_pdf_rejects_scanned_like_deck():
    # Pages with no text content at all (simulating a scanned/image-only deck)
    doc = fitz.open()
    doc.new_page()
    doc.new_page()
    buf = io.BytesIO()
    doc.save(buf)

    with pytest.raises(DeckExtractionError, match="scan|image"):
        extract_pdf(buf.getvalue(), "scanned.pdf")


def test_extract_pptx_happy_path():
    pptx_bytes = make_pptx_bytes(
        [
            "Market Size: $10B TAM",
            "Team: 3 founders, 20 years combined experience",
        ]
    )
    deck = extract_pptx(pptx_bytes, "deck.pptx")

    assert deck.source_format == "pptx"
    assert deck.slide_count == 2
    assert "Market Size" in deck.slides[0].text
    assert "Team" in deck.slides[1].text


def test_extract_pptx_rejects_corrupted_file():
    with pytest.raises(DeckExtractionError):
        extract_pptx(b"not a real pptx", "broken.pptx")


def test_extract_pptx_rejects_empty_deck():
    prs = Presentation()
    buf = io.BytesIO()
    prs.save(buf)

    with pytest.raises(DeckExtractionError):
        extract_pptx(buf.getvalue(), "empty.pptx")


def test_deck_to_prompt_text_labels_slides():
    pdf_bytes = make_pdf_bytes(
        [
            "Slide one content with enough text to clear the scanned-deck heuristic.",
            "Slide two content with enough text to clear the scanned-deck heuristic.",
        ]
    )
    deck = extract_pdf(pdf_bytes, "deck.pdf")
    prompt_text = deck.to_prompt_text()

    assert "--- Slide 1 ---" in prompt_text
    assert "--- Slide 2 ---" in prompt_text
    assert "Slide one content" in prompt_text
